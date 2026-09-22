"""Prepare and verify a relocatable Windows CPython runtime without modifying sources.

The output is optional local analysis data, never an npm package or a venv.
Only fetch-assets uses the network; prepare/verify never download models.
"""
from __future__ import annotations

import argparse
import csv
from email.parser import Parser
import hashlib
import importlib
import importlib.metadata
import json
import os
from pathlib import Path, PurePosixPath, PureWindowsPath
import posixpath
import shutil
import stat
import subprocess
import sys
import tempfile
import urllib.request
import uuid
import zipfile

PACKAGE = Path(__file__).resolve().parents[1]
REVISION = '12af15fef9d0ac838c3f475bfbbf26d2060dd4f5'
REPOSITORY = 'models--m-a-p--MERT-v1-95M'
HEAD_SHA = 'deaceb291f7974deb688167d3639b7f6eb66eb0715824668618393777d1e07a5'
BTC_SHA = '1673d23f8f9a55ae7f9e8b80a51da616debb22675b8d8b67ea6ce0ef37b0ab51'
DISTUTILS_SHIM = "import os; var = 'SETUPTOOLS_USE_DISTUTILS'; enabled = os.environ.get(var, 'local') == 'local'; enabled and __import__('_distutils_hack').add_shim();"
# Recorded once per distribution in the manifest; see the manifest-level `provenance`.
LOCAL_SOURCE = 'local validated environment'
PINS = {'torch': '2.3.1+cpu', 'torchaudio': '2.3.1+cpu', 'transformers': '4.44.0'}
IMPORTS = ['torch', 'torchaudio', 'transformers', 'librosa', 'music21', 'mir_eval',
           'pretty_midi', 'numpy', 'yaml', 'pytorch_lightning', 'sklearn', 'requests',
           'PIL', 'docx', 'lxml.etree']


def digest(path: Path) -> str:
    with path.open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def new_output(path: Path) -> None:
    """Claim a new directory exclusively; an existing directory is not resumable."""
    path.mkdir(exist_ok=False)


def site_pth_location(destination: str) -> bool:
    """Whether `site` would read a `.pth` file at this destination.

    `site.addsitedir` scans `*.pth` in the directory itself and nowhere else, so a
    `.pth` inside a package is ordinary data: torchmetrics ships pickled `.pth`
    state dicts under `torchmetrics/functional/image/lpips_models/`.
    """
    return PurePosixPath(destination).parent == PurePosixPath('Lib/site-packages')


def keep_pth(name: str, content: bytes) -> bool:
    """Remove only the reviewed setuptools shim; reject other executable .pth files."""
    try:
        text = content.decode('utf-8-sig')
    except UnicodeDecodeError as error:
        raise ValueError(f'Undecodable .pth requires review: {name}') from error
    if name == 'distutils-precedence.pth' and text.strip() == DISTUTILS_SHIM:
        return False
    for line in text.splitlines():
        line = line.strip()
        if not line or line.startswith('#'):
            continue
        if line.startswith(('import ', 'import\t')):
            raise ValueError(f'Executable .pth requires review: {name}')
        if PureWindowsPath(line).drive or line.startswith(('/', '\\')) or '..' in PureWindowsPath(line).parts:
            raise ValueError(f'External .pth path refused: {name}')
    return True


def record_destination(value: str) -> str | None:
    """Map an installed RECORD member into the standalone prefix, excluding launchers."""
    if '\\' in value or PureWindowsPath(value).drive or value.startswith('/'):
        raise ValueError(f'Absolute RECORD path: {value}')
    result = posixpath.normpath(posixpath.join('Lib/site-packages', value))
    if result.startswith('../') or result == '..':
        raise ValueError(f'RECORD escapes runtime: {value}')
    if result.startswith('Scripts/') or '__pycache__' in PurePosixPath(result).parts or result.endswith('.pyc'):
        return None
    if not result.startswith(('Lib/site-packages/', 'Library/', 'share/', 'Include/')):
        raise ValueError(f'Unsupported RECORD resource: {value}')
    return result


def regular(path: Path) -> None:
    info = path.lstat()
    if not stat.S_ISREG(info.st_mode) or getattr(info, 'st_file_attributes', 0) & stat.FILE_ATTRIBUTE_REPARSE_POINT:
        raise ValueError(f'Non-regular source file: {path}')


def copy_file(source: Path, target: Path, materialize_within: Path | None = None) -> None:
    if materialize_within is not None:
        source = source.resolve(strict=True)
        if not source.is_relative_to(materialize_within.resolve()):
            raise ValueError('Model cache link escapes its repository')
    regular(source)
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        if digest(target) != digest(source):
            raise ValueError(f'Conflicting source files for {target}')
        return
    with source.open('rb') as incoming, target.open('xb') as outgoing:
        shutil.copyfileobj(incoming, outgoing, 1024 * 1024)


def file_record(path: Path, root: Path) -> dict:
    regular(path)
    return {'path': path.relative_to(root).as_posix(), 'bytes': path.stat().st_size, 'sha256': digest(path)}


def inventory(root: Path) -> list[dict]:
    return [file_record(path, root) for path in sorted(root.rglob('*')) if path.is_file()
            and path.name != 'runtime-manifest.json']


def verify_files(root: Path) -> dict:
    """Reject changed, missing, extra or escaping payload members."""
    manifest = json.loads((root / 'runtime-manifest.json').read_text(encoding='utf-8'))
    if manifest.get('version') != 1:
        raise ValueError('Unknown runtime manifest version')
    expected = set()
    for item in manifest['files']:
        name = item['path']
        if '\\' in name or PureWindowsPath(name).drive or name.startswith('/') or '..' in PurePosixPath(name).parts:
            raise ValueError('Unsafe runtime manifest path')
        if name in expected:
            raise ValueError('Duplicate runtime manifest path')
        expected.add(name)
        path = root / name
        if not path.resolve().is_relative_to(root.resolve()) or file_record(path, root) != item:
            raise ValueError(f'Runtime file changed: {name}')
    actual = {item['path'] for item in inventory(root)}
    if actual != expected:
        raise ValueError(f'Runtime inventory changed: extra={sorted(actual - expected)}, missing={sorted(expected - actual)}')
    return manifest


def assets() -> list[dict]:
    return json.loads((Path(__file__).with_name('runtime-assets.json')).read_text(encoding='utf-8'))


def fetch_assets(directory: Path) -> None:
    new_output(directory)
    for item in assets():
        target = directory / item['filename']
        with urllib.request.urlopen(item['url'], timeout=60) as response, target.open('xb') as stream:
            total = 0
            while chunk := response.read(1024 * 1024):
                total += len(chunk)
                if total > 8 * 1024 * 1024:
                    raise ValueError('Supplementary artifact exceeds byte limit')
                stream.write(chunk)
        if digest(target) != item['sha256']:
            target.unlink()
            raise ValueError(f'Artifact hash mismatch: {item["filename"]}')


def prepare(args: argparse.Namespace) -> None:
    if os.name != 'nt':
        raise ValueError('This preparer targets Windows CPython 3.11 x64 only')
    base, venv, output = args.python_base.resolve(), args.venv.resolve(), args.output.resolve()
    if output.is_relative_to(base) or output.is_relative_to(venv):
        raise ValueError('Output must not be inside a source tree')
    if output.exists():
        raise FileExistsError(
            f'Refusing existing output: {output}. An interrupted build leaves PREPARING.json; '
            'remove that directory to rebuild')
    site = venv / 'Lib/site-packages'
    distributions = []
    for metadata in sorted(site.glob('*.dist-info/METADATA')):
        header = Parser().parsestr(metadata.read_text(encoding='utf-8'))
        distributions.append({'name': header['Name'], 'version': header['Version'],
                              'source': LOCAL_SOURCE, 'metadata_sha256': digest(metadata)})
    versions = {row['name'].lower().replace('_', '-'): row['version'] for row in distributions}
    if any(versions.get(name) != version for name, version in PINS.items()):
        raise ValueError('Source environment does not match the verified CPU dependency pins')
    for path, expected in ((args.head, HEAD_SHA), (args.btc, BTC_SHA)):
        if digest(path) != expected:
            raise ValueError(f'Checkpoint hash mismatch: {path.name}')
    for item in assets():
        if digest(args.assets / item['filename']) != item['sha256']:
            raise ValueError(f'Artifact hash mismatch: {item["filename"]}')
    new_output(output)
    marker = output / 'PREPARING.json'
    marker.write_text('{"complete":false}\n', encoding='utf-8')
    runtime = output / 'python'
    omitted = []
    # The standalone base supplies all stdlib/native resources, never its pip site-packages or launchers.
    for current, directories, filenames in os.walk(base):
        folder = Path(current)
        for name in list(directories):
            relative = (folder / name).relative_to(base).as_posix()
            if name == '__pycache__' or relative in ('Lib/site-packages', 'Scripts'):
                directories.remove(name)
            elif (folder / name).lstat().st_file_attributes & stat.FILE_ATTRIBUTE_REPARSE_POINT:
                raise ValueError('Standalone base contains a directory link')
        for name in filenames:
            if name.endswith('.pyc'):
                continue
            if name == 'pyvenv.cfg' or name.endswith('._pth'):
                raise ValueError('Standalone base already contains environment/path overrides')
            source = folder / name
            copy_file(source, runtime / source.relative_to(base))
    for record in sorted(site.glob('*.dist-info/RECORD')):
        with record.open(encoding='utf-8', newline='') as stream:
            for row in csv.reader(stream):
                destination = record_destination(row[0])
                if destination is None:
                    continue
                source = (site / row[0]).resolve(strict=True)
                if not source.is_relative_to(venv):
                    raise ValueError('Distribution resource escapes the source venv')
                if site_pth_location(destination) and source.suffix == '.pth' and not keep_pth(source.name, source.read_bytes()):
                    omitted.append({'path': destination, 'sha256': digest(source), 'reason': 'reviewed setuptools startup shim'})
                    continue
                if source.name in ('sitecustomize.py', 'usercustomize.py') or source.suffix == '.egg-link':
                    raise ValueError('Custom interpreter startup requires explicit review')
                copy_file(source, runtime / destination)
    # Fixed wheels supplement absent skill dependencies; never run pip or replace an installed distribution.
    for item in assets():
        path = args.assets / item['filename']
        if path.suffix != '.whl':
            copy_file(path, output / 'licenses' / path.name)
            continue
        with zipfile.ZipFile(path) as archive:
            for member in archive.infolist():
                if member.is_dir():
                    continue
                if '..' in PurePosixPath(member.filename).parts or PureWindowsPath(member.filename).drive or '\\' in member.filename or member.filename.startswith('/'):
                    raise ValueError('Unsafe wheel member')
                if '.data/' in member.filename:
                    raise ValueError('Wheel uses unsupported installation scheme')
                target = runtime / 'Lib/site-packages' / member.filename
                target.parent.mkdir(parents=True, exist_ok=True)
                with target.open('xb') as stream:
                    stream.write(archive.read(member))
    (runtime / 'python311._pth').write_text('.\nDLLs\nLib\nLib/site-packages\nimport site\n', encoding='utf-8')
    model_repo = args.hf_home / 'hub' / REPOSITORY
    snapshot = model_repo / 'snapshots' / REVISION
    for name in ('config.json', 'preprocessor_config.json', 'configuration_MERT.py', 'modeling_MERT.py', 'pytorch_model.bin'):
        copy_file(snapshot / name, output / 'models/hf-cache/hub' / REPOSITORY / 'snapshots' / REVISION / name,
                  materialize_within=model_repo)
    # huggingface_hub recreates this cache marker on first use, so a cache copied without
    # it is not equal to the validated one and the first read stops being read-only.
    copy_file(args.hf_home / 'hub/version.txt', output / 'models/hf-cache/hub/version.txt',
              materialize_within=args.hf_home)
    refs = output / 'models/hf-cache/hub' / REPOSITORY / 'refs'
    refs.mkdir()
    (refs / 'main').write_text(REVISION, encoding='utf-8')
    copy_file(args.head, output / 'models/J_all.ckpt')
    copy_file(args.btc, output / 'models/data/btc_model_large_voca.pt')
    for path in sorted((PACKAGE / 'python/data').iterdir()):
        if path.suffix in ('.json', '.yaml', '.npy'):
            copy_file(path, output / 'models/data' / path.name)
    copy_file(PACKAGE / 'python/LICENSE-Music2Emo', output / 'licenses/LICENSE-Music2Emo')
    copy_file(PACKAGE / 'SOURCES.md', output / 'licenses/SOURCES.md')
    (output / 'licenses/MERT-ATTRIBUTION.txt').write_text(
        'm-a-p/MERT-v1-95M\nRevision: ' + REVISION + '\n'
        'https://huggingface.co/m-a-p/MERT-v1-95M\n'
        'CC-BY-NC-4.0: non-commercial use only. No model weights were downloaded.\n'
        'The local snapshot contains no model card; retain this attribution and the full license.\n', encoding='utf-8')
    for metadata in sorted((runtime / 'Lib/site-packages').glob('*.dist-info/METADATA')):
        header = Parser().parsestr(metadata.read_text(encoding='utf-8'))
        if header['Name'].lower().replace('_', '-') in ('python-docx', 'lxml'):
            distributions.append({'name': header['Name'], 'version': header['Version'],
                                  'source': 'hash-pinned PyPI wheel', 'metadata_sha256': digest(metadata)})
    marker.unlink()
    manifest = {'version': 1, 'python': {'version': '3.11.16', 'build': (base / 'BUILD').read_text().strip(),
                'source': 'existing uv-managed standalone CPython; upstream archive URL not recorded'},
                'models': {'mert_repository': 'm-a-p/MERT-v1-95M', 'revision': REVISION, 'license': 'CC-BY-NC-4.0'},
                'distributions': distributions, 'supplementary_assets': assets(), 'omitted': omitted,
                'provenance': f'Distributions marked "{LOCAL_SOURCE}" were copied byte-for-byte from the locally '
                              'validated CPython 3.11.16 environment; nothing was downloaded, resolved or upgraded '
                              'during preparation. Each metadata_sha256 identifies the recorded name and version, and '
                              'files[] identifies the installed bytes. Upstream wheel hashes are not recorded, so this '
                              'manifest proves the installed bytes, not their PyPI origin.',
                'policy': 'No venv launchers, bytecode, executable startup hooks, credentials or external path entries.',
                'files': inventory(output)}
    with (output / 'runtime-manifest.json').open('x', encoding='utf-8') as stream:
        json.dump(manifest, stream, ensure_ascii=False, indent=2)
        stream.write('\n')
    print(json.dumps({'output': str(output), 'files': len(manifest['files']),
                      'bytes': sum(row['bytes'] for row in manifest['files'])}))


def isolated_env(root: Path, work: Path) -> dict[str, str]:
    system = os.environ['SystemRoot']
    python = root / 'python'
    work.mkdir(exist_ok=True)
    temp = work / 'temp'
    temp.mkdir(exist_ok=True)
    return {'SystemRoot': system, 'WINDIR': system,
            'PATH': os.pathsep.join(map(str, [python, python / 'DLLs', python / 'Library/bin', Path(system) / 'System32'])),
            'USERPROFILE': str(work), 'HOME': str(work), 'APPDATA': str(work), 'LOCALAPPDATA': str(work),
            'TEMP': str(temp), 'TMP': str(temp), 'OMP_NUM_THREADS': '1', 'MKL_NUM_THREADS': '1',
            'OPENBLAS_NUM_THREADS': '1', 'HF_HUB_OFFLINE': '1', 'TRANSFORMERS_OFFLINE': '1',
            'HF_HUB_DISABLE_TELEMETRY': '1', 'HF_HOME': str(root / 'models/hf-cache'),
            'HF_MODULES_CACHE': str(work / 'hf-modules'), 'MPLCONFIGDIR': str(work / 'matplotlib'),
            'NUMBA_CACHE_DIR': str(work / 'numba'), 'TORCH_HOME': str(work / 'torch')}


def probe(root: Path, work: Path, audio: Path | None) -> None:
    prefix = (root / 'python').resolve()
    if sys.version_info[:3] != (3, 11, 16) or Path(sys.prefix).resolve() != prefix or Path(sys.base_prefix).resolve() != prefix:
        raise ValueError('Interpreter version/prefix is not the relocated standalone runtime')
    if not sys.flags.isolated or not sys.flags.dont_write_bytecode:
        raise ValueError('Probe must run with -I -B')
    if any(not Path(item).resolve().is_relative_to(prefix) for item in sys.path):
        raise ValueError(f'External sys.path: {sys.path}')
    imported = {}
    for name in IMPORTS:
        module = importlib.import_module(name)
        location = Path(module.__file__).resolve()
        if not location.is_relative_to(prefix):
            raise ValueError(f'Imported {name} outside runtime')
        imported[name] = str(location)
    versions = {name: importlib.metadata.version(name) for name in PINS}
    if versions != PINS:
        raise ValueError(f'CPU dependency versions differ: {versions}')
    # Read/write a document and image, not just import their package names.
    from io import BytesIO
    from docx import Document
    from PIL import Image
    document = Document()
    document.add_paragraph('portable runtime')
    memory = BytesIO()
    document.save(memory)
    memory.seek(0)
    assert Document(memory).paragraphs[0].text == 'portable runtime'
    image = BytesIO()
    Image.new('RGB', (2, 2)).save(image, format='PNG')
    image.seek(0)
    assert Image.open(image).size == (2, 2)
    if audio:
        import torchaudio
        rate = torchaudio.info(str(audio)).sample_rate
        waveform, rate = torchaudio.load(str(audio), num_frames=12 * rate)
        torchaudio.save(str(work / 'inspect.wav'), waveform, rate)
    print(json.dumps({'python': sys.version, 'prefix': sys.prefix, 'base_prefix': sys.base_prefix,
                      'path': sys.path, 'imports': imported, 'versions': versions,
                      'skill_roundtrips': ['docx', 'Pillow']}))


def run_probe(root: Path, work: Path, audio: Path | None = None) -> dict:
    argv = [str(root / 'python/python.exe'), '-I', '-B', '-X', 'utf8', str(Path(__file__).resolve()),
            'probe', '--root', str(root), '--work', str(work)]
    if audio:
        argv += ['--audio', str(audio)]
    result = subprocess.run(argv, cwd=work, env=isolated_env(root, work), capture_output=True, text=True,
                            encoding='utf-8', timeout=180)
    if result.returncode:
        raise RuntimeError(result.stderr + result.stdout)
    return json.loads(result.stdout.splitlines()[-1])


def verify(args: argparse.Namespace) -> None:
    root = args.root.resolve()
    manifest = verify_files(root)
    report = {'files_verified': len(manifest['files']), 'bytes_verified': sum(row['bytes'] for row in manifest['files'])}
    with tempfile.TemporaryDirectory(prefix='verify-', dir=root.parent) as temporary:
        work = Path(temporary)
        report['original'] = run_probe(root, work, args.audio)
        if args.relocate:
            relocated = root.with_name(root.name + '-relocated-' + uuid.uuid4().hex)
            root.rename(relocated)
            try:
                report['relocated'] = run_probe(relocated, work)
                if root.exists():
                    raise RuntimeError('Original location unexpectedly exists during relocation check')
            finally:
                if root.exists():
                    raise RuntimeError('Refusing to overwrite the occupied original runtime location')
                relocated.rename(root)
        if args.audio:
            request = {'id': 'portable-inspect', 'method': 'analyse', 'params': {
                'audio_path': str(work / 'inspect.wav'), 'data_dir': str(root / 'models/data'),
                'weights_path': str(root / 'models/J_all.ckpt')}}
            result = subprocess.run([str(root / 'python/python.exe'), '-I', '-B', '-X', 'utf8',
                                     str(PACKAGE / 'python/worker_main.py')], cwd=work,
                                    input=json.dumps(request) + '\n', capture_output=True, text=True, encoding='utf-8',
                                    timeout=600, env=isolated_env(root, work))
            if result.returncode:
                raise RuntimeError(result.stderr + result.stdout)
            messages = [json.loads(line) for line in result.stdout.splitlines() if line.strip()]
            answer = next(row for row in messages if row['id'] == 'portable-inspect')
            if answer['status'] != 'ok':
                raise RuntimeError(json.dumps(answer) + result.stderr)
            analysis = answer['result']
            if not (1 <= analysis['valence'] <= 9 and 1 <= analysis['arousal'] <= 9):
                raise ValueError('Inspection returned invalid measurements')
            report['inspect'] = analysis
            report['scratch_cleaned'] = not list((work / 'temp').glob('dsh-bgm-*'))
            if not report['scratch_cleaned']:
                raise ValueError('Inference scratch was not cleaned')
    verify_files(root)
    with args.report.open('x', encoding='utf-8') as stream:
        json.dump(report, stream, ensure_ascii=False, indent=2)
        stream.write('\n')
    print(json.dumps({'report': str(args.report), 'files_verified': report['files_verified'],
                      'relocated': 'relocated' in report, 'inspect': 'inspect' in report}))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    fetch = commands.add_parser('fetch-assets')
    fetch.add_argument('--output', type=Path, required=True)
    build = commands.add_parser('prepare')
    for name in ('python-base', 'venv', 'hf-home', 'head', 'btc', 'assets', 'output'):
        build.add_argument('--' + name, type=Path, required=True)
    check = commands.add_parser('verify')
    check.add_argument('--root', type=Path, required=True)
    check.add_argument('--report', type=Path, required=True)
    check.add_argument('--relocate', action='store_true')
    check.add_argument('--audio', type=Path)
    child = commands.add_parser('probe')
    child.add_argument('--root', type=Path, required=True)
    child.add_argument('--work', type=Path, required=True)
    child.add_argument('--audio', type=Path)
    args = parser.parse_args()
    if args.command == 'fetch-assets':
        fetch_assets(args.output)
    elif args.command == 'prepare':
        prepare(args)
    elif args.command == 'verify':
        verify(args)
    else:
        probe(args.root, args.work, args.audio)


if __name__ == '__main__':
    main()
