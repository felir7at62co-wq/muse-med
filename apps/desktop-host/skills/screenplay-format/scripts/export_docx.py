"""Convert ordered screenplay Markdown files and verify all saved paragraphs."""
import io
import os
import tempfile
import json
import re
import sys
import zipfile
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def export(request):
    sources = []
    total = 0
    pinned = request.get("contents")
    if pinned is not None and len(pinned) != len(request["inputs"]):
        raise ValueError("Pinned content count differs from input count")
    for index, name in enumerate(request["inputs"]):
        data = pinned[index].encode("utf-8") if pinned is not None else Path(name).read_bytes()
        total += len(data)
        if total > request["max_input_bytes"]:
            raise ValueError("Markdown inputs exceed the export size limit")
        sources.append(data.decode("utf-8-sig"))
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.5)
    normal = document.styles["Normal"]
    normal.font.name = "SimSun"
    normal.font.size = Pt(12)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    expected = []
    for source in sources:
        for line in source.splitlines():
            if not line.strip():
                continue
            heading = re.match(r"^#{1,6}\s+(.+?)\s*$", line)
            content = heading.group(1) if heading else line
            parts = re.split(r"(\*\*[^*]+\*\*)", content)
            plain = "".join(part[2:-2] if part.startswith("**") and part.endswith("**") else part for part in parts)
            episode = bool(re.fullmatch(r"第[0-9一二三四五六七八九十百零〇]+集", plain.strip()))
            scene = bool(re.match(r"^\d+-\d+\s+", plain))
            paragraph = document.add_paragraph()
            for part in parts:
                bold = part.startswith("**") and part.endswith("**") and len(part) > 4
                run = paragraph.add_run(part[2:-2] if bold else part)
                run.bold = bold or episode or scene or bool(heading)
                if episode:
                    run.font.size = Pt(16)
            paragraph.paragraph_format.keep_with_next = episode or scene or bool(heading) or plain.startswith("人物：")
            paragraph.paragraph_format.widow_control = True
            if episode and expected:
                paragraph.paragraph_format.page_break_before = True
            expected.append(plain)
    if not expected:
        raise ValueError("No screenplay content to export")
    if request.get("verify_only"):
        with zipfile.ZipFile(request["output"]) as archive:
            if sum(item.file_size for item in archive.infolist()) > request["max_docx_bytes"]:
                raise ValueError("Word contents exceed the verification size limit")
        reopened = Document(request["output"])
        body = reopened.element.body
        if any(child.tag not in (qn("w:p"), qn("w:sectPr")) for child in body):
            raise ValueError("Unexpected Word body content outside screenplay paragraphs")
        if [p.text for p in reopened.paragraphs] != expected:
            raise ValueError("Word paragraph content differs from accepted Markdown")
        for paragraph in reopened.paragraphs:
            inline = "".join(node.text or "" if node.tag == qn("w:t") else "\t" if node.tag == qn("w:tab") else "\n"
                             for node in paragraph._p.iter() if node.tag in (qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr")))
            if inline != paragraph.text:
                raise ValueError("Unexpected nested text in a screenplay paragraph")
        for tag in ("w:drawing", "w:pict", "w:object", "w:instrText", "w:footnoteReference", "w:endnoteReference"):
            if next(body.iter(qn(tag)), None) is not None:
                raise ValueError("Unexpected embedded body content")
        for saved_section in reopened.sections:
            for extra in (saved_section.header, saved_section.first_page_header, saved_section.even_page_header,
                          saved_section.footer, saved_section.first_page_footer, saved_section.even_page_footer):
                if any(node.text for node in extra._element.iter(qn("w:t"))):
                    raise ValueError("Unexpected screenplay text outside the accepted body")
                for tag in ("w:drawing", "w:pict", "w:object", "w:instrText"):
                    if next(extra._element.iter(qn(tag)), None) is not None:
                        raise ValueError("Unexpected embedded header or footer content")
                if any(field.get(qn("w:instr"), "").strip() != "PAGE" for field in extra._element.iter(qn("w:fldSimple"))):
                    raise ValueError("Only page numbering is allowed outside the accepted body")
        print(json.dumps({"paragraphs": len(expected)}, ensure_ascii=False))
        return
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    reopened = Document(buffer)
    if [p.text for p in reopened.paragraphs] != expected:
        raise ValueError("Word paragraph content differs from Markdown")
    output = Path(request["output"])
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=".screenplay-", suffix=".docx", dir=output.parent)
    try:
        with os.fdopen(descriptor, "wb") as target:
            target.write(buffer.getvalue())
        os.link(temporary, output)
    finally:
        os.unlink(temporary)
    print(json.dumps({"paragraphs": len(expected)}, ensure_ascii=False))


if __name__ == "__main__":
    export(json.load(sys.stdin))
