"""Exercise packaged Python, native media libraries and document output without network or user media."""
import importlib
import importlib.util
import importlib.metadata
import json
import os
from pathlib import Path
import socket
import subprocess
import sys
import tempfile


def main():
    """Check relocatable imports, dependency declarations, encode/subtitles, draft probing and Word round trips."""
    root = Path(sys.argv[1]).resolve()
    expected = root / 'python'
    assert Path(sys.prefix).resolve() == expected
    assert Path(sys.base_prefix).resolve() == expected
    assert not (expected / 'pyvenv.cfg').exists()
    for path in sys.path:
        if path:
            assert Path(path).resolve().is_relative_to(expected), path
    def deny_network(*args, **kwargs):
        raise AssertionError('Media smoke must not use network')
    socket.create_connection = deny_network
    os.environ['HF_HUB_OFFLINE'] = '1'
    os.environ['HF_HUB_DISABLE_TELEMETRY'] = '1'
    assert not (root / 'models').exists()
    for name in ['faster_whisper', 'ctranslate2', 'av', 'onnxruntime', 'tokenizers', 'huggingface_hub']:
        assert importlib.util.find_spec(name) is None, name
    from packaging.requirements import Requirement
    for dist in importlib.metadata.distributions():
        for declared in dist.requires or []:
            req = Requirement(declared)
            if req.marker and not req.marker.evaluate({'extra': ''}):
                continue
            actual = importlib.metadata.version(req.name)
            assert actual in req.specifier, (dist.metadata['Name'], declared, actual)
    for name in ['docx', 'pydub', 'pymediainfo', 'pyJianYingDraft', 'requests', 'PIL', 'zhconv', 'cv2', 'scipy']:
        module = importlib.import_module(name)
        assert Path(module.__file__).resolve().is_relative_to(expected), name
    from zhconv import convert
    assert convert('繁體字幕', 'zh-cn') == '繁体字幕'
    ffmpeg = str(root / 'ffmpeg' / 'bin' / 'ffmpeg.exe')
    ffprobe = str(root / 'ffmpeg' / 'bin' / 'ffprobe.exe')
    with tempfile.TemporaryDirectory(prefix='muse-media-smoke-') as temporary:
        work = Path(temporary)
        (work / 'captions.srt').write_text('1\n00:00:00,000 --> 00:00:01,000\nMedia smoke\n', encoding='utf-8')
        subprocess.run([ffmpeg, '-hide_banner', '-loglevel', 'error', '-y', '-f', 'lavfi', '-i', 'color=c=black:s=320x240:r=10:d=1', '-f', 'lavfi', '-i', 'anullsrc=r=16000:cl=mono', '-t', '1', '-vf', 'subtitles=captions.srt', '-af', 'loudnorm', '-c:v', 'libx264', '-threads', '1', '-c:a', 'aac', 'smoke.mp4'], cwd=work, check=True)
        probe = subprocess.run([ffprobe, '-v', 'error', '-show_streams', '-of', 'json', str(work / 'smoke.mp4')], capture_output=True, text=True, check=True)
        assert {s['codec_type'] for s in json.loads(probe.stdout)['streams']} == {'video', 'audio'}
        subprocess.run([sys.executable, '-B', str(Path(__file__).with_name('smoke-cjk-font.py')),
                        '--ffmpeg', ffmpeg, '--fonts', str(root / 'fonts'), '--output', str(work / 'cjk')], check=True)
        import cv2
        from scipy.optimize import linear_sum_assignment
        video = cv2.VideoCapture(str(work / 'smoke.mp4'))
        try:
            ok, frame = video.read()
            assert ok and frame.shape[:2] == (240, 320)
        finally:
            video.release()
        rows, columns = linear_sum_assignment([[5, 1], [1, 5]])
        assert columns.tolist() == [1, 0]
        from pymediainfo import MediaInfo
        assert any(t.track_type == 'Video' for t in MediaInfo.parse(str(work / 'smoke.mp4')).tracks)
        import pyJianYingDraft as draft
        material = draft.VideoMaterial(str(work / 'smoke.mp4'))
        assert material.duration > 0
        from docx import Document
        document = Document()
        document.add_paragraph('Media smoke')
        document.save(work / 'smoke.docx')
        assert Document(work / 'smoke.docx').paragraphs[0].text == 'Media smoke'
    print('media smoke: isolated imports, locked dependencies, FFmpeg/AAC/H264/subtitles/loudnorm, MediaInfo/draft and Word round trip passed')


if __name__ == '__main__':
    main()
